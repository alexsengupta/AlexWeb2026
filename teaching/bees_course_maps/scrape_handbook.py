"""
UNSW Handbook scraper for the BEES course maps
==============================================

Reads every course code mentioned in the program workbooks (the Structure and
Streams sheets of each file listed in programs.xlsx), fetches each course page
from the UNSW Handbook, and writes the results to handbook.xlsx.

    python scrape_handbook.py              # latest handbook year that responds
    python scrape_handbook.py --year 2027  # a specific year
    python scrape_handbook.py --dry-run    # list the codes and stop

What it writes
    handbook.xlsx
        Info      - handbook year, date scraped, course count   (rewritten)
        Courses   - one row per course                          (rewritten)
        Overrides - hand corrections                            (never touched)

The Overrides sheet is where you fix things by hand (a course that is not in
the handbook yet, a wrong title, a missing term). Any non-blank cell there
replaces the scraped value in the app; blank cells fall back to the scrape.

How it works
    The handbook is a Next.js app: every course page embeds its data as JSON
    in a <script id="__NEXT_DATA__"> tag, so no API key is needed.

Requirements
    pip install requests beautifulsoup4 lxml openpyxl
"""
from __future__ import annotations

import argparse
import datetime as dt
import os
import re
import sys
import time

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
INDEX_FILE = os.path.join(HERE, "programs.xlsx")
HANDBOOK_FILE = os.path.join(HERE, "handbook.xlsx")
CODE_RE = re.compile(r"\b[A-Z]{4}\d{4}\b")
PAUSE_SECONDS = 0.5  # be polite to the handbook server

COLUMNS = ["Code", "Title", "UoC", "Terms", "Prerequisites", "Prerequisite codes", "Exclusions",
           "Overview", "School", "Faculty", "Delivery", "Contact hours", "Handbook URL", "Error"]
WIDTHS = [10, 42, 6, 22, 50, 30, 20, 60, 30, 20, 12, 12, 50, 30]


# ─────────────────────────────────────────────────────────────────────────────
# 1. Which codes to fetch: everything mentioned in the program workbooks
# ─────────────────────────────────────────────────────────────────────────────
def program_files() -> list[str]:
    wb = load_workbook(INDEX_FILE, read_only=True, data_only=True)
    ws = wb["Programs"]
    files = []
    for row in ws.iter_rows(values_only=True):
        if not row or not row[0]:
            continue
        name = str(row[0]).strip()
        if name.lower() == "file" or not name.lower().endswith(".xlsx"):
            continue
        files.append(name)
    return files


def codes_in_workbook(path: str) -> set[str]:
    codes: set[str] = set()
    wb = load_workbook(path, read_only=True, data_only=True)
    # Only course columns are scanned, so codes mentioned in notes are ignored.
    wanted = {"Structure": ("Course", "Alternative 1", "Alternative 2", "Alternative 3"),
              "Streams": None}  # None = every "Course N" column
    for sheet, cols in wanted.items():
        if sheet not in wb.sheetnames:
            continue
        rows = wb[sheet].iter_rows(values_only=True)
        header = None
        for row in rows:
            if header is None:
                if row and any(str(v).strip() in ("Course", "Stream") for v in row if v):
                    header = [str(v).strip() if v else "" for v in row]
                continue
            for i, v in enumerate(row):
                h = header[i] if i < len(header) else ""
                ok = h.startswith("Course") if cols is None else h in cols
                if ok and isinstance(v, str):
                    codes.update(CODE_RE.findall(v))
    return codes


def all_codes() -> list[str]:
    codes: set[str] = set()
    for f in program_files():
        path = os.path.join(HERE, f)
        if not os.path.exists(path):
            print(f"  WARNING: {f} listed in programs.xlsx but not found", file=sys.stderr)
            continue
        found = codes_in_workbook(path)
        print(f"  {f}: {len(found)} codes")
        codes |= found
    return sorted(codes)


# ─────────────────────────────────────────────────────────────────────────────
# 2. Fetching
# ─────────────────────────────────────────────────────────────────────────────
def level(code: str) -> str:
    return "postgraduate" if int(code[4]) >= 5 else "undergraduate"


def url_for(code: str, year: int) -> str:
    return f"https://www.handbook.unsw.edu.au/{level(code)}/courses/{year}/{code}"


def strip_html(html: str) -> str:
    return BeautifulSoup(html or "", "lxml").get_text(separator=" ").strip() if html else ""


def first(obj, *keys, default=""):
    for k in keys:
        if isinstance(obj, list):
            obj = obj[0] if obj else None
        if isinstance(obj, dict):
            obj = obj.get(k)
        else:
            return default
    return str(obj).strip() if obj not in (None, "") else default


def detect_latest_year(session: requests.Session, probe_code: str, start: int) -> int:
    """Try next year, this year, last year with a known course; first 200 wins."""
    for y in (start + 1, start, start - 1):
        try:
            r = session.get(url_for(probe_code, y), timeout=20)
            if r.status_code == 200 and '"pageType":"ErrorPage"' not in r.text:
                return y
        except requests.RequestException:
            pass
    return start


def fetch_course(session: requests.Session, code: str, year: int) -> dict:
    url = url_for(code, year)
    out = {c: "" for c in COLUMNS}
    out["Code"], out["Handbook URL"] = code, url
    try:
        resp = session.get(url, timeout=20)
    except requests.RequestException as e:
        out["Error"] = f"request failed: {e}"
        return out
    if resp.status_code == 404:
        out["Error"] = f"not in {year} handbook"
        return out
    if resp.status_code != 200:
        out["Error"] = f"HTTP {resp.status_code}"
        return out

    tag = BeautifulSoup(resp.text, "lxml").find("script", id="__NEXT_DATA__")
    if not tag:
        out["Error"] = "no __NEXT_DATA__ in page"
        return out
    try:
        import json
        props = json.loads(tag.string).get("props", {}).get("pageProps", {})
    except Exception as e:
        out["Error"] = f"JSON error: {e}"
        return out
    if props.get("pageType") == "ErrorPage":
        out["Error"] = f"not in {year} handbook"
        return out
    pc = props.get("pageContent") or {}
    if not pc.get("cl_code"):
        out["Error"] = "pageContent empty"
        return out

    rules = "; ".join(strip_html(r.get("description", "")) for r in (pc.get("enrolment_rules") or []) if r.get("description"))
    exclusions = ", ".join(first(e, "associated_ai_cl_id", "value") for e in (pc.get("exclusion") or []) if first(e, "associated_ai_cl_id", "value"))
    dv = (pc.get("hb_delivery_variations") or [{}])[0]
    od = pc.get("offering_detail") or {}

    out.update({
        "Title": pc.get("title", ""),
        "UoC": pc.get("credit_points", ""),
        "Terms": od.get("offering_terms", ""),
        "Prerequisites": rules,
        "Prerequisite codes": ", ".join(sorted(set(CODE_RE.findall(rules)))),
        "Exclusions": exclusions,
        "Overview": strip_html(pc.get("description", "")),
        "School": first(pc.get("school_detail"), "name"),
        "Faculty": first(pc.get("school_detail"), "parent", "value"),
        "Delivery": first(dv.get("delivery_mode"), "value"),
        "Contact hours": dv.get("contact_hours", ""),
    })
    return out


# ─────────────────────────────────────────────────────────────────────────────
# 3. Writing handbook.xlsx (keeping Overrides intact)
# ─────────────────────────────────────────────────────────────────────────────
def write_handbook(rows: list[dict], year: int) -> None:
    override_rows: list[list] = []
    if os.path.exists(HANDBOOK_FILE):
        old = load_workbook(HANDBOOK_FILE)
        if "Overrides" in old.sheetnames:
            override_rows = [list(r) for r in old["Overrides"].iter_rows(values_only=True)]

    wb = Workbook()
    wb.remove(wb.active)
    head_fill, head_font, note_font = PatternFill("solid", fgColor="1F3A5F"), Font(bold=True, color="FFFFFF"), Font(italic=True, color="666666")

    def add_sheet(name, header, data, note, widths):
        ws = wb.create_sheet(name)
        ws.cell(row=1, column=1, value=note).font = note_font
        for c, h in enumerate(header, 1):
            cell = ws.cell(row=2, column=c, value=h)
            cell.fill, cell.font = head_fill, head_font
        for r, row in enumerate(data, 3):
            for c, v in enumerate(row, 1):
                ws.cell(row=r, column=c, value=v)
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.freeze_panes = "A3"

    add_sheet("Info", ["Setting", "Value"], [
        ["Handbook year", year],
        ["Scraped on", dt.date.today().isoformat()],
        ["Courses", len(rows)],
        ["Source", "scrape_handbook.py"],
    ], "Written by scrape_handbook.py. Do not edit.", [16, 60])
    add_sheet("Courses", COLUMNS, [[r[c] for c in COLUMNS] for r in rows],
              "Written by scrape_handbook.py. Do not edit; put corrections in the Overrides sheet.", WIDTHS)

    # Preserve the Overrides sheet exactly (its own note line and header included)
    if override_rows:
        ws = wb.create_sheet("Overrides")
        for r, row in enumerate(override_rows, 1):
            for c, v in enumerate(row, 1):
                ws.cell(row=r, column=c, value=v)
        ws.cell(row=1, column=1).font = note_font
        for c in range(1, len(COLUMNS) + 1):
            ws.cell(row=2, column=c).fill, ws.cell(row=2, column=c).font = head_fill, head_font
        for c, w in enumerate(WIDTHS, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.freeze_panes = "A3"
    else:
        add_sheet("Overrides", COLUMNS, [],
                  "Hand corrections. Any non-blank cell here replaces the scraped value for that course. Survives re-scraping.", WIDTHS)

    wb.save(HANDBOOK_FILE)


# ─────────────────────────────────────────────────────────────────────────────
def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--year", type=int, help="handbook year (default: latest that responds)")
    ap.add_argument("--dry-run", action="store_true", help="list the codes and exit")
    args = ap.parse_args()

    print("Collecting course codes from program workbooks:")
    codes = all_codes()
    print(f"  {len(codes)} unique codes\n")
    if args.dry_run:
        print(", ".join(codes))
        return

    session = requests.Session()
    session.headers["User-Agent"] = "Mozilla/5.0 (BEES course map scraper; contact a.sengupta@unsw.edu.au)"

    year = args.year or detect_latest_year(session, "MATH1131", dt.date.today().year)
    print(f"Scraping the {year} handbook")

    rows = []
    for i, code in enumerate(codes, 1):
        row = fetch_course(session, code, year)
        status = row["Error"] or f"{row['Title']}  [{row['Terms'] or 'no terms listed'}]"
        print(f"  [{i:2d}/{len(codes)}] {code}: {status}")
        rows.append(row)
        time.sleep(PAUSE_SECONDS)

    write_handbook(rows, year)
    missing = [r["Code"] for r in rows if r["Error"]]
    print(f"\nWrote {HANDBOOK_FILE}")
    if missing:
        print(f"{len(missing)} course(s) not found: {', '.join(missing)}")
        print("Add titles for these in the Overrides sheet if you want them to display properly.")


if __name__ == "__main__":
    main()
